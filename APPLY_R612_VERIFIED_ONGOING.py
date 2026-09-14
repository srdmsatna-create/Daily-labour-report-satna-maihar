from pathlib import Path
import json, re, sys
from openpyxl import load_workbook

ROOTS = [
    Path.cwd(),
    Path.home() / "Downloads",
    Path.home() / "Desktop",
    Path.home() / "Documents",
]
NAMES = [
    "SRDM_SATNA_R6.12_ONGOING_FINAL_PART1_13-09-2026.xlsx",
    "SRDM_SATNA_R6.12_ONGOING_FINAL_PART2_13-09-2026.xlsx",
    "SRDM_SATNA_R6.12_ONGOING_FINAL_PART3_13-09-2026.xlsx",
]

def find_file(name):
    for root in ROOTS:
        p = root / name
        if p.exists():
            return p
    for root in ROOTS[1:]:
        if root.exists():
            try:
                for p in root.rglob(name):
                    return p
            except Exception:
                pass
    return None

parts = []
for name in NAMES:
    p = find_file(name)
    if not p:
        raise SystemExit("ERROR: File not found: " + name + "\nPlease keep the 3 R6.12 FINAL PART files in Downloads.")
    parts.append(p)

print("Using verified files:")
for p in parts:
    print(" ", p)

rows = []
for p in parts:
    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    hdr = [c.value for c in ws[3]]
    ix = {h:i for i,h in enumerate(hdr)}
    need = ['S.No.','District','Janpad','Gram Panchayat','Sub Engineer / Upyantri','Cluster',
            'Work FY','Work Status','Work Code','Work Name','Final Work Category','Original Work Type',
            'Total Sanction (Rs)','Booked Wages (Rs)','Booked Material (Rs)','Total Booked (Rs)',
            'Expenditure %','Total Mandays','Current FY Mandays']
    miss = [x for x in need if x not in ix]
    if miss:
        raise SystemExit(f"ERROR: Missing columns in {p.name}: {miss}")
    for r in ws.iter_rows(min_row=4, values_only=True):
        code = str(r[ix['Work Code']] or '').strip()
        if not code:
            continue
        cat = str(r[ix['Final Work Category']] or '').strip()
        if cat == 'IAY Houses':
            cat = 'PMAY-G'
        rows.append({
            'sno': r[ix['S.No.']],
            'district': str(r[ix['District']] or '').strip(),
            'janpad': str(r[ix['Janpad']] or '').strip(),
            'panchayat': str(r[ix['Gram Panchayat']] or '').strip(),
            'engineer': str(r[ix['Sub Engineer / Upyantri']] or '').strip(),
            'cluster': str(r[ix['Cluster']] or '').strip(),
            'fy': str(r[ix['Work FY']] or '').strip(),
            'status': str(r[ix['Work Status']] or '').strip(),
            'code': code,
            'name': str(r[ix['Work Name']] or '').strip(),
            'type': str(r[ix['Original Work Type']] or '').strip(),
            'finalCategory': cat,
            'sanction': float(r[ix['Total Sanction (Rs)']] or 0),
            'bookedWage': float(r[ix['Booked Wages (Rs)']] or 0),
            'bookedMaterial': float(r[ix['Booked Material (Rs)']] or 0),
            'booked': float(r[ix['Total Booked (Rs)']] or 0),
            'expPct': float(r[ix['Expenditure %']] or 0),
            'mandays': float(r[ix['Total Mandays']] or 0),
            'currentFYMandays': float(r[ix['Current FY Mandays']] or 0),
        })

codes = {r['code'] for r in rows}
if len(rows) != 15687 or len(codes) != 15687:
    raise SystemExit(f"ERROR: R6.12 validation failed: rows={len(rows)}, unique={len(codes)}; expected 15687.")

cats = {}
for r in rows:
    cats[r['finalCategory']] = cats.get(r['finalCategory'], 0) + 1

if cats.get('Ek Bagiya') != 755:
    raise SystemExit(f"ERROR: Ek Bagiya ongoing = {cats.get('Ek Bagiya')}; expected 755.")
if cats.get('PMAY-G') != 10807:
    raise SystemExit(f"ERROR: PMAY-G including IAY = {cats.get('PMAY-G')}; expected 10807.")
if cats.get('IAY Houses', 0) != 0:
    raise SystemExit("ERROR: IAY Houses still separate.")

outp = Path('ongoing-details.js')
oldmap = {}
if outp.exists():
    try:
        t = outp.read_text(encoding='utf-8-sig').strip()
        prefix = 'window.ONGOING_DETAILS='
        if t.startswith(prefix):
            b = t[len(prefix):]
            if b.endswith(';'): b = b[:-1]
            old = json.loads(b)
            oldmap = {str(x.get('code','')).strip():x for x in old}
    except Exception as e:
        print("WARNING: could not read old ongoing-details.js:", e)

merged = []
for r in rows:
    old = oldmap.get(r['code'], {})
    x = dict(old)
    x.update(r)
    for k in ('recoveryDone','recoveryWork','recoveryAmount','recoveryWorkCount',
              'nregaAprJunMandays','julyMandays'):
        if k in old:
            x[k] = old[k]
    merged.append(x)

outp.write_text('window.ONGOING_DETAILS=' + json.dumps(merged, ensure_ascii=False, separators=(',',':')) + ';\n',
                 encoding='utf-8')

ip = Path('index.html')
patched = False
if ip.exists():
    s = ip.read_text(encoding='utf-8-sig')
    old = "function canonicalCategoryName(v){const s=clean(v);if(/^(?:pmay\\s*-?\\s*g|pmayg)(?:\\s+houses?)?$/i.test(s))return 'PMAY-G';"
    new = "function canonicalCategoryName(v){const s=clean(v);if(/^(?:pmay\\s*-?\\s*g|pmayg)(?:\\s+houses?)?$/i.test(s)||/^iay(?:\\s+houses?)?$/i.test(s))return 'PMAY-G';"
    if old in s:
        s = s.replace(old, new, 1)
        patched = True
    elif "function canonicalCategoryName(v){const s=clean(v);" in s and "/^iay" not in s:
        s = s.replace("function canonicalCategoryName(v){const s=clean(v);",
                      "function canonicalCategoryName(v){const s=clean(v);if(/^iay(?:\\s+houses?)?$/i.test(s))return 'PMAY-G';", 1)
        patched = True
    ip.write_text(s, encoding='utf-8')

print()
print("SUCCESS DATA VALIDATION")
print("R6.12 ongoing works       :", len(merged))
print("PMAY-G (PMAY + IAY)       :", cats.get('PMAY-G'))
print("IAY separate category     : 0")
print("Ek Bagiya ongoing         :", cats.get('Ek Bagiya'))
print("Ek Bagiya verified total  : 756 (755 ongoing + 1 physically completed)")
print("Index category patch      :", "OK" if patched else "Already OK / not needed")
